"""Открытые данные(T-26…T-31): импортёр xlsx/csv, провенанс записей, объекты, источники, импорт через API,
отсутствие риска и сигналов у записей источников, разделение с вымышленным демо-набором(ТЗ §7)."""
import csv
import io
import json
from pathlib import Path

import openpyxl
import pytest

from app import importer
from app.risk_model import SHARED

XLSX = SHARED.parent / "docs" / "Усть-Каменогорск_объекты.xlsx"
need_xlsx = pytest.mark.skipif(not XLSX.is_file(), reason="нет docs/Усть-Каменогорск_объекты.xlsx")
AUTHOR = "Нестерова А."
PROVENANCE = ("source", "sourceName", "sourceUrl", "sourceRecordId", "collectedAt")
RISK_KEYS = {"score", "zone", "needsCheck", "factors", "risk", "signal", "signals", "hits", "observations", "summary"}
JUDGEMENT = ("наруш", "незакон", "нелегал", "мошен", "обман", "риск", "зона", "балл")


def _rows():
    """Строки листа как есть: заголовок набора, пояснение, пустая строка, шапка, данные."""
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _csv(rows, delimiter=";", bom=True) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return (("﻿" if bom else "") + buf.getvalue()).encode("utf-8")


def _xlsx(rows) -> bytes:
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(c, data: bytes, name="data.xlsx", ctype="application/octet-stream", **form):
    return c.post("/api/sources/import", files={"file": (name, data, ctype)}, data=form)


@pytest.fixture(scope="module")
def parsed():
    return importer.parse(XLSX.read_bytes(), XLSX.name) if XLSX.is_file() else None


# ---------- парсер ----------

@need_xlsx
def test_parse_counts_and_objects(parsed):
    p = parsed
    assert (p.format, p.read, len(p.records), p.skipped, p.duplicates, p.errors) == ("xlsx", 103, 103, 0, 0, [])
    assert {s: sum(r["source"] == s for r in p.records) for s in importer.SOURCE_ORDER} == {
        "homeportal": 16, "homeportal_section": 1, "ashyq": 86}
    assert p.preamble[1].startswith("103 записей источников")
    objs = importer.group(p.records)
    assert len(objs) == 15 and sum(o["records"] for o in objs) == 103
    assert sorted(o["key"] for o in objs if o["lat"] is None) == ["ab-residence-2", "diamond", "prestizh-sparkle-diamond"]
    zh = next(o for o in objs if o["key"] == "zhanym")  # блоки и очереди одного ЖК - один объект
    assert (zh["name"], zh["records"], zh["sources"], zh["fundingMethods"]) == (
        'ЖК "Жаным"', 32, ["homeportal", "ashyq"], ["guarantee"])
    assert (zh["lat"], zh["lng"], zh["developerBin"]) == (49.921112, 82.617788, "171240023488")
    assert {o["key"] for o in objs} >= {"zolotoi-kvadrat", "tumar-towers-oskemen", "asyltas", "asylym", "orda", "park-line"}
    # координаты источника далеко от города - оставлены как есть, но помечены
    assert next(o for o in objs if o["key"] == "renesans")["flags"] == {"coordsFarKm": 188}
    assert any("Строка 8: координаты" in w and "оставлены как в источнике" in w for w in p.warnings)


@need_xlsx
def test_parse_values(parsed):
    by = {r["sourceRecordId"]: r for r in parsed.records}
    hp, kz = by["1505"], by["4828"]
    # оба формата дат → ISO; время сбора - UTC; дата загрузки у поставщика - как у поставщика, без пояса
    assert (hp["docDate"], hp["start"], hp["plannedEnd"]) == ("2026-10-07", "2026-09-18", "2027-11-18")
    assert (kz["start"], kz["plannedEnd"]) == ("2024-06-19", "2025-06-19")
    assert hp["collectedAt"] == "2026-10-15T11:18:00+00:00" and kz["collectedAt"] == "2026-10-15T11:17:00+00:00"
    assert (hp["providerLoadedAt"], kz["providerLoadedAt"]) == (None, "2026-06-15T17:35:09")
    # БИН - 12 цифр строкой(с ведущим нулём), у Казреестра - нет
    assert (hp["developerBin"], hp["agentBin"], by["1487"]["developerBin"]) == ("220740026935", "190740008285",
                                                                                 "050640012251")
    assert kz["developerBin"] is None and sum(r["developerBin"] is not None for r in parsed.records) == 17
    # "0" в координатах - нет координат; координаты есть у 15 записей
    assert (by["1487"]["lat"], by["1487"]["lng"]) == (None, None) and (hp["lat"], hp["lng"]) == (49.917561, 82.619097)
    assert sum(r["lat"] is not None for r in parsed.records) == 15
    # способы привлечения: код + исходный текст
    assert {(r["fundingText"], r["fundingMethod"]) for r in parsed.records if r["fundingText"]} == {
        ("Разрешение МИО", "mio"), ('Гарантия АО "КЖК"', "guarantee"), ("Участие БВУ", "bvu")}
    assert kz["fundingText"] is None and kz["fundingMethod"] is None
    # коды источников
    assert {(r["sourceName"], r["source"]) for r in parsed.records} == {
        ("Homeportal — карточка ЖК", "homeportal"), ("Казреестр — CSV Ashyq Data", "ashyq"),
        ("Homeportal — отдельный раздел", "homeportal_section")}
    # пустые ячейки - null, готовность и период отчёта - числом и "ГГГГ-ММ"
    assert (hp["actNumber"], hp["actDate"], hp["readiness"], hp["site"][:8]) == (None, None, None, "https://")
    rep = [r for r in parsed.records if r["readiness"] is not None]
    assert sorted(r["readiness"] for r in rep) == [42.41, 62.39] and {r["reportPeriod"] for r in rep} == {"2026-08"}
    # провенанс - у каждой записи
    assert all(r.get(f) for r in parsed.records for f in PROVENANCE) and all(r["row"] >= 5 for r in parsed.records)
    assert len({(r["source"], r["sourceRecordId"]) for r in parsed.records}) == 103


def test_value_normalizers():
    assert importer._date("07.10.2026") == ("2026-10-07", None) and importer._date("2024-06-19") == ("2024-06-19", None)
    assert importer._date("31.02.2026")[0] is None and importer._date(None) == (None, None)
    assert importer._utc("15.10.2026 11:17 UTC") == ("2026-10-15T11:17:00+00:00", None)
    assert importer._utc("15.10.2026 16:17 +05:00")[0] == "2026-10-15T11:17:00+00:00"
    assert importer._utc("2026-10-15T11:17:00Z")[0] == "2026-10-15T11:17:00+00:00"
    assert importer._local("2026-06-15 17:35:09") == ("2026-06-15T17:35:09", None)
    assert importer._bin(50640012251) == ("050640012251", None) and importer._bin(" 220740026935 ")[0] == "220740026935"
    assert importer._bin("12345")[0] is None and importer._bin("") == (None, None)
    assert importer._number("49,917561") == 49.917561 and importer._number("abc") is None
    assert importer._period("08.2026") == ("2026-08", None) and importer._percent("62,39%") == (62.39, None)
    assert importer._percent("120")[0] is None
    assert [importer.funding_code(t) for t in ("Гарантия АО \"КЖК\"", "Разрешение МИО", "Участие БВУ", "", "иное")] == [
        "guarantee", "mio", "bvu", None, None]
    assert importer.source_code("Homeportal — отдельный раздел") == "homeportal_section"
    assert importer.source_code("Homeportal", "https://homeportal.kz/ru/Illegal-projects/1813") == "homeportal_section"
    assert importer.source_code("Казреестр — CSV Ashyq Data") == "ashyq" and importer.source_code("Krisha") is None
    # нормализованное название и ключ объекта
    names = ['ЖК "Жаным", Блок А', 'ЖК "Жаным" (блок А)', "ЖК «Жаным»", "жк жаным, очередь 2", "МЖК Жаным, Пятно 3"]
    assert {importer.object_key(n) for n in names} == {"zhanym"}
    assert importer.object_key("ЖК «Престиж» (Sparkle Diamond) (ГС: Инвестор)") == "prestizh-sparkle-diamond"
    assert importer.object_key("Tumar Towers Oskemen, Здание 1") == importer.object_key('ЖК "Tumar Towers Oskemen"')
    assert importer.object_key('ЖК "AB-Residence 2"') == "ab-residence-2"
    assert importer.org_name('Товарищество с ограниченной ответственностью "НАУАН Корп"') == importer.org_name(
        'ТОО "НАУАН Корп"') == "науан корп"


@need_xlsx
def test_registry_json_matches_xlsx(parsed):
    """shared/registry_ust_kamenogorsk.json собран импортёром из docs-xlsx: пометка происхождения и те же записи."""
    data = json.loads(importer.REGISTRY.read_text(encoding="utf-8"))
    meta = data["meta"]
    assert meta["origin"] == "source" and meta["builtFrom"] == "docs/Усть-Каменогорск_объекты.xlsx"
    assert meta["builtAt"].endswith("+00:00") and meta["collectedAt"] == "2026-10-15T11:24:00+00:00"
    assert "не выводы системы" in data["_note"] and "не рассчитываются" in data["_note"]
    assert meta["report"]["read"] == 103 and meta["report"]["objects"] == 15 and meta["report"]["withoutCoords"] == 3
    assert [s["code"] for s in data["sources"]] == ["homeportal", "homeportal_section", "ashyq"]
    strip = lambda r: {k: v for k, v in r.items() if k != "origin"}  # noqa: E731
    assert data["records"] == [strip(r) for r in parsed.records]


def test_registry_json_objects_match_api(client):
    """Готовые объекты в JSON реестра(офлайн-режим фронтенда) - ровно то, что отдаёт GET /api/registry после сида."""
    data = importer.load_registry()
    assert (len(data["objects"]), len(data["withoutCoords"])) == (12, 3)
    assert {"objects": data["objects"], "withoutCoords": data["withoutCoords"]} == client.get("/api/registry").json()
    assert importer.split(importer.group(data["records"])) == {"objects": data["objects"],
                                                               "withoutCoords": data["withoutCoords"]}
    renesans = next(o for o in data["objects"] if o["key"] == "renesans")
    ab = next(o for o in data["withoutCoords"] if o["key"] == "ab-residence-2")
    assert renesans["flags"] == {"coordsFarKm": 188} and ab["flags"] == {"sourceSection": "Без разрешения на дольщиков"}


@need_xlsx
def test_cli_dry_run(capsys, tmp_path):
    out = tmp_path / "reg.json"
    assert importer.main([str(XLSX), "--out", str(out)]) == 0
    built = json.loads(out.read_text(encoding="utf-8"))
    assert built["meta"]["report"]["accepted"] == 103 and (len(built["objects"]), len(built["withoutCoords"])) == (12, 3)
    assert {k: built[k] for k in ("records", "objects", "withoutCoords")} == {
        k: importer.load_registry()[k] for k in ("records", "objects", "withoutCoords")}  # файл в shared/ - актуальный
    assert importer.main([str(XLSX), "--dry-run", "--out", str(tmp_path / "none.json")]) == 0
    assert not (tmp_path / "none.json").exists() and "объектов 15, из них без координат 3" in capsys.readouterr().out
    bad = tmp_path / "x.txt"
    bad.write_text("x")
    assert importer.main([str(bad), "--dry-run"]) == 2


# ---------- API: источники и объекты ----------

def test_sources(client):
    r = client.get("/api/sources")
    assert r.status_code == 200
    by = {s["code"]: s for s in r.json()}
    assert [s["code"] for s in r.json()] == ["homeportal", "homeportal_section", "ashyq", "demo"]
    assert (by["homeportal"]["records"], by["homeportal_section"]["records"], by["ashyq"]["records"]) == (16, 1, 86)
    assert by["ashyq"]["providerLoadedAt"] == "2026-06-15T17:35:09" and by["ashyq"]["url"].startswith("https://ashyq")
    assert by["homeportal"]["collectedAt"] == "2026-10-15T11:24:00+00:00" and by["homeportal"]["objects"] == 13
    demo = by["demo"]
    assert demo["name"] == "Демонстрационный набор (вымышленные ЖК)" and demo["records"] == 14
    assert demo["counts"] == {"complexes": 14, "developers": 7, "ads": 12, "appeals": 5} and demo["url"] is None
    assert all("counts" not in s for s in r.json()[:3])


def test_registry_list_and_item(client):
    d = client.get("/api/registry").json()
    assert (len(d["objects"]), len(d["withoutCoords"])) == (12, 3)
    assert all(o["lat"] is not None and o["lng"] is not None for o in d["objects"])
    assert all(o["lat"] is None for o in d["withoutCoords"])
    keys = {"key", "origin", "name", "address", "developer", "developerBin", "agent", "fundingMethods", "stage",
            "plannedEnd", "lat", "lng", "records", "sources", "collectedAt", "flags"}
    for o in d["objects"] + d["withoutCoords"]:
        assert set(o) == keys and o["origin"] == "source" and o["collectedAt"].endswith("+00:00")
    assert client.get("/api/registry", params={"withCoords": "true"}).json()["withoutCoords"] == []
    assert client.get("/api/registry", params={"withCoords": "false"}).json()["objects"] == []
    # поиск: название(в т. ч. латиницей ↔ кириллицей), адрес, застройщик, БИН, номер документа
    q = lambda s: sorted(o["key"] for part in client.get("/api/registry", params={"q": s}).json().values()  # noqa: E731
                         for o in part)
    assert q("жаным") == q("Zhanym") == q("«ЖАНЫМ»") == ["zhanym"]
    assert q("Асылтас") == ["asyltas"] and q("Mercury") == ["park-line", "renesans"] == q("220740026935")
    assert q("ДПГ-25-16-078/194") == ["zolotoi-kvadrat"] and q("Кайсенова") == ["zolotoi-kvadrat"]
    assert q("нет такого объекта") == [] and len(q("  ")) == 15

    r = client.get("/api/registry/zhanym")
    assert r.status_code == 200
    item = r.json()
    assert item["object"]["key"] == "zhanym" and len(item["records"]) == 32
    assert [x["source"] for x in item["records"]][:3] == ["homeportal"] * 3
    for rec in item["records"]:
        assert rec["origin"] == "source" and rec["key"] == "zhanym" and all(rec[f] for f in PROVENANCE)
        assert rec["importedAt"].endswith("+00:00") and rec["importFile"] == "Усть-Каменогорск_объекты.xlsx"
    assert client.get("/api/registry/nope").status_code == 404
    assert client.get("/api/registry/nope").json()["detail"] == "Объект открытых данных не найден"


def test_no_risk_for_source_records(client):
    """T-31: у записей источников нет риска, зоны, "нужна проверка"; сигнал и риск по ключу реестра - 404."""
    d = client.get("/api/registry").json()
    for o in d["objects"] + d["withoutCoords"]:
        assert not RISK_KEYS & set(o) and not RISK_KEYS & set(o["flags"])
        item = client.get(f"/api/registry/{o['key']}").json()
        assert all(not RISK_KEYS & set(rec) for rec in item["records"])
    for key in ("zhanym", "ab-residence-2"):
        r = client.post("/api/signals", json={"complexId": key, "author": AUTHOR})
        assert r.status_code == 404 and "не рассчитываются" in r.json()["detail"]
        assert client.get(f"/api/complexes/{key}/risk").status_code == 404
        assert client.get(f"/api/complexes/{key}").status_code == 404
    assert client.post("/api/signals", json={"complexId": "nope", "author": AUTHOR}).json()["detail"] == "ЖК не найден"
    assert client.get("/api/signals", params={"complexId": "zhanym"}).json() == []
    risk_ids = {r["complexId"] for r in client.get("/api/risk").json()}
    assert risk_ids == {c["id"] for c in client.get("/api/complexes").json()} and "zhanym" not in risk_ids
    # ИИ сопоставляет публикации только с ЖК демо-реестра
    hyps = client.post("/api/analyze", json={"text": "ЖК «Жаным», Есенберлина 92: бронь 100 000 ₸"}).json()["hypotheses"]
    assert all(h["complexId"].startswith("c") for h in hyps)


def test_source_section_is_attributed_fact(client):
    """Запись из раздела Homeportal - только факт источника: раздел, ссылка, дата сбора; без оценочных слов."""
    item = client.get("/api/registry/ab-residence-2").json()
    o, (rec,) = item["object"], item["records"]
    assert o["flags"] == {"sourceSection": "Без разрешения на дольщиков"} and o["sources"] == ["homeportal_section"]
    assert (rec["source"], rec["sourceSection"], rec["sourceUrl"], rec["collectedAt"]) == (
        "homeportal_section", "Без разрешения на дольщиков", "https://homeportal.kz/ru/Illegal-projects/1813",
        "2026-10-15T11:18:00+00:00")
    text = json.dumps({**o, "records": [{k: v for k, v in rec.items() if k not in ("sourceUrl", "note")}]},
                      ensure_ascii=False).lower()
    assert not any(w in text for w in JUDGEMENT)


def test_demo_does_not_match_open_data(client):
    """T-31(ТЗ §7 "Безопасность"): вымышленные ЖК и застройщики не совпадают с записями открытых данных -
    сценарии нарушений не могут быть прочитаны как обвинение реальной компании."""
    data = importer.load_registry()
    recs = data["records"]
    names = {importer.object_name(r["name"]) for r in recs} | {importer.object_key(r["name"]) for r in recs}
    skeletons = {importer.skeleton_text(importer.object_name(r["name"])) for r in recs}
    for c in client.get("/api/complexes").json():
        n = importer.object_name(c["name"])
        assert n not in names and importer.object_key(c["name"]) not in names, c["name"]
        assert importer.skeleton_text(n) not in skeletons, c["name"]
        assert c["origin"] == "demo"
    orgs = {importer.org_name(r.get(f)) for r in recs for f in ("developer", "agent") if r.get(f)}
    bins = {r.get(f) for r in recs for f in ("developerBin", "agentBin") if r.get(f)}
    for d in client.get("/api/developers").json():
        assert importer.org_name(d["name"]) not in orgs and d["bin"] not in bins, d["name"]
        assert d["origin"] == "demo"


def test_origin_in_outputs(client):
    assert {c["origin"] for c in client.get("/api/complexes").json()} == {"demo"}
    assert client.get("/api/complexes/c2").json()["origin"] == "demo"
    assert {d["origin"] for d in client.get("/api/developers").json()} == {"demo"}
    assert {a["origin"] for a in client.get("/api/appeals").json()} == {"demo"}
    assert {a["origin"] for a in client.get("/api/ads").json()} == {"demo"}
    r = client.post("/api/ads", json={"platform": "telegram", "account": "@x", "text": "Квартиры в ЖК «Алтай Резиденс»",
                                      "author": AUTHOR})
    assert r.status_code == 201 and r.json()["origin"] == "analyst"
    assert next(a for a in client.get("/api/ads").json() if a["id"] == r.json()["id"])["origin"] == "analyst"
    # провенанс оснований по документу проходит в карточку ЖК как есть
    srcs = client.get("/api/complexes/c2").json()["basis"]["sources"]
    assert set(srcs) == {"guarantee", "mio", "bvu"} and srcs["guarantee"]["url"] == "https://homeportal.kz"


# ---------- импорт через API ----------

@need_xlsx
def test_import_same_file_dedup_and_journal(client):
    r = _import(client, XLSX.read_bytes(), XLSX.name, author=AUTHOR)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert (rep["file"], rep["format"], rep["read"], rep["created"], rep["updated"], rep["unchanged"], rep["skipped"],
            rep["duplicates"]) == (XLSX.name, "xlsx", 103, 0, 0, 103, 0, 0)
    assert (rep["records"], rep["objects"], rep["withoutCoords"]) == (103, 15, 3) and rep["errors"] == []
    assert rep["sources"] == {"homeportal": 16, "homeportal_section": 1, "ashyq": 86}
    log = client.get("/api/audit", params={"kind": "data"}).json()
    e = next(x for x in log if x["id"] == rep["eventId"])
    assert (e["who"], e["kind"], e["complexId"]) == (AUTHOR, "data", None)
    assert e["what"] == f"Импорт: 0 новых, 0 обновлено, 103 без изменений, 0 пропущено — {XLSX.name}"
    # исходная загрузка из сида - тоже в журнале
    assert any(x["who"] == "Система" and x["what"].startswith("Загружены сведения открытых источников: 103 записи")
               for x in log)
    # без автора - "Система"; автор параметром запроса
    assert client.get("/api/audit", params={"kind": "data"}).json()[-1]["who"] == AUTHOR
    _import(client, XLSX.read_bytes(), XLSX.name)
    assert client.get("/api/audit", params={"kind": "data"}).json()[-1]["who"] == "Система"
    client.post("/api/sources/import?author=Исаев%20Т.", files={"file": (XLSX.name, XLSX.read_bytes(), "")})
    assert client.get("/api/audit", params={"kind": "data"}).json()[-1]["who"] == "Исаев Т."


@need_xlsx
@pytest.mark.parametrize("delimiter,bom", [(";", True), (",", True), (",", False)])
def test_import_csv(client, delimiter, bom):
    rows = _rows()
    hi = next(i for i, r in enumerate(rows) if r[0] == "Источник")
    data = rows[:hi + 1] + rows[hi + 1:hi + 11]
    data[hi + 1] = list(data[hi + 1])
    data[hi + 1][24] = "999001"  # новая запись
    data[hi + 2] = list(data[hi + 2])
    data[hi + 2][2] = "г. Усть-Каменогорск, новый адрес, 1"  # изменилось поле существующей записи
    r = _import(client, _csv(data, delimiter, bom), "part.csv", "text/csv")
    assert r.status_code == 200, r.text
    rep = r.json()
    assert (rep["format"], rep["read"], rep["created"], rep["updated"], rep["unchanged"]) == ("csv", 10, 1, 1, 8)
    assert rep["records"] == 104
    new = client.get("/api/registry/park-line").json()["records"]
    assert [x["sourceRecordId"] for x in new] == ["1505", "999001"] and new[1]["importFile"] == "part.csv"
    prestizh = client.get("/api/registry/prestizh-sparkle-diamond").json()["records"][0]
    assert prestizh["address"] == "г. Усть-Каменогорск, новый адрес, 1" and prestizh["lat"] is None
    # те же значения из CSV(строки, "0" в координатах) нормализуются так же, как из xlsx
    assert new[0]["lat"] == 49.917561 and new[0]["docDate"] == "2026-10-07" and new[0]["developerBin"] == "220740026935"
    r = _import(client, _csv(data, delimiter, bom), "part.csv", "text/csv")
    assert (r.json()["created"], r.json()["updated"], r.json()["unchanged"]) == (0, 0, 10)


def test_import_skips_and_duplicates(client):
    head = [h for h, _, _ in importer.COLUMNS]
    row = lambda src, rid, name="ЖК «Тест»": [src, name] + [None] * 22 + [rid, None]  # noqa: E731
    rows = [["Набор"], [], head, row("Казреестр — CSV Ashyq Data", "t1"), row("Казреестр — CSV Ashyq Data", "t1"),
            row("Krisha", "t2"), row("Homeportal — карточка ЖК", None), row("Homeportal — карточка ЖК", "t3", None),
            [None] * 26]
    r = _import(client, _xlsx(rows), "mix.xlsx")
    rep = r.json()
    assert r.status_code == 200 and (rep["read"], rep["created"], rep["skipped"], rep["duplicates"]) == (5, 1, 3, 1)
    assert any("Строка 6: неизвестный источник «Krisha»" in e for e in rep["errors"])
    assert any("Строка 7: нет «ID в источнике»" in e for e in rep["errors"])
    assert any("Строка 8: нет «ЖК / объект»" in e for e in rep["errors"])
    assert any("Строка 5: повтор записи ashyq #t1" in w for w in rep["warnings"])
    e = client.get("/api/audit", params={"kind": "data"}).json()[-1]
    assert e["what"] == "Импорт: 1 новых, 0 обновлено, 0 без изменений, 4 пропущено — mix.xlsx" and e["who"] == "Система"
    t = client.get("/api/registry/test").json()["records"][0]
    assert (t["source"], t["collectedAt"], t["lat"]) == ("ashyq", None, None)


def test_import_errors(client, monkeypatch):
    big = b"x" * (importer.MAX_BYTES + 1)
    r = _import(client, big, "big.csv")
    assert r.status_code == 413 and r.json()["detail"] == importer.TOO_BIG
    for name, ctype in (("data.txt", "text/plain"), ("pic.png", "image/png"), ("noext", "application/octet-stream")):
        r = _import(client, b"abc", name, ctype)
        assert r.status_code == 415 and "xlsx" in r.json()["detail"], name
    assert "xls" in _import(client, b"abc", "old.xls").json()["detail"] and _import(client, b"a", "old.xls").status_code == 415
    cases = [(b"not a zip", "broken.xlsx", "повреждён"), (b"", "empty.csv", "пустой"),
             ("Источник;ЖК\nx;y\n".encode("cp1251"), "cp.csv", "UTF-8"),
             ("a;b;c\n1;2;3\n".encode(), "nohead.csv", "заголовка"),
             (_xlsx([["a", "b"], [1, 2]]), "nohead.xlsx", "заголовка"), (b"PK\x03\x04zzz", "zip.csv", "Excel")]
    for data, name, word in cases:
        r = _import(client, data, name)
        assert r.status_code == 422 and word in r.json()["detail"], (name, r.text)
    # CSV без расширения, но с типом text/csv - тоже CSV; zip-бомба - 422
    assert _import(client, b"a;b\n", "noext", "text/csv").status_code == 422
    monkeypatch.setattr(importer, "MAX_UNZIPPED", 10)
    r = _import(client, _xlsx([["a"]]), "bomb.xlsx")
    assert r.status_code == 422 and "после распаковки" in r.json()["detail"]
    assert client.post("/api/sources/import").status_code == 422  # нет файла
    assert client.get("/api/audit", params={"kind": "data"}).json()[-1]["what"].startswith("Загружены сведения")


def test_import_limit_before_reading_body():
    """Content-Length больше предела - 413 до чтения тела(как у OCR)."""
    import asyncio

    from app.main import UploadLimit

    sent, reads = [], []

    async def app_(scope, receive, send):
        raise AssertionError("обработчик не должен вызываться")

    async def receive():
        reads.append(1)
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(message):
        sent.append(message)

    limit = UploadLimit.LIMITS["/api/sources/import"][0]
    scope = {"type": "http", "method": "POST", "path": "/api/sources/import",
             "headers": [(b"content-length", str(limit + 1).encode())]}
    asyncio.run(UploadLimit(app_)(scope, receive, send))
    assert sent[0]["status"] == 413 and json.loads(sent[1]["body"])["detail"] == importer.TOO_BIG and not reads


def test_seed_without_registry_file(client, monkeypatch, tmp_path):
    """Нет файла реестра(например, в другой сборке) - демо-набор работает, открытых данных просто нет."""
    from app import seed

    monkeypatch.setattr(importer, "REGISTRY", tmp_path / "none.json")
    seed.reset()
    assert client.get("/api/registry").json() == {"objects": [], "withoutCoords": []}
    assert [s["code"] for s in client.get("/api/sources").json()] == ["demo"]
    assert len(client.get("/api/complexes").json()) == 14


def test_registry_file_is_in_shared():
    """Docker-образ видит только shared/ - реестр должен лежать там(seed грузит его в source_records)."""
    assert importer.REGISTRY.parent == SHARED and importer.REGISTRY.is_file()
    assert Path(importer.REGISTRY).name == "registry_ust_kamenogorsk.json"
